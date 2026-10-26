"""Geração da planilha de controle (.xlsx) dos pedidos.

São duas abas:
  • "Aprovados"  — só o que está confirmado, é a lista de produção.
  • "Todos"      — tudo, menos os recusados.

Pedidos recusados NUNCA entram na planilha: eles não viraram venda, só
poluiriam a conta.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config import EXPORTS_DIR
from app.models import APROVADO, ENTREGUE, RECUSADO, Pedido
from app.whatsapp import telefone_formatado

PLANILHA_PADRAO = EXPORTS_DIR / "pedidos.xlsx"

COLUNAS = [
    ("Pedido", 9),
    ("Situação", 20),
    ("Cliente", 24),
    ("Telefone", 18),
    ("Data da entrega", 16),
    ("Dia da semana", 15),
    ("Como receber", 20),
    ("Endereço", 40),
    ("Itens", 38),
    ("Total (R$)", 12),
    ("Observações", 30),
    ("Pedido feito em", 18),
]

DIAS = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]

_CABECALHO_FILL = PatternFill("solid", fgColor="17346B")  # azul-marinho da logo
_CABECALHO_FONTE = Font(bold=True, color="FBC81C", size=11)  # amarelo da logo
_LINHA_FINA = Side(style="thin", color="D4D9E2")
_BORDA = Border(left=_LINHA_FINA, right=_LINHA_FINA, top=_LINHA_FINA, bottom=_LINHA_FINA)


def _escrever_aba(ws, pedidos: list[Pedido], vazio: str) -> None:
    for idx, (titulo, largura) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=1, column=idx, value=titulo)
        celula.fill = _CABECALHO_FILL
        celula.font = _CABECALHO_FONTE
        celula.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = largura
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"

    if not pedidos:
        ws.cell(row=2, column=1, value=vazio).font = Font(italic=True, color="7A7A7A")
        return

    # Os mais próximos de entregar primeiro — é assim que ela vai usar a planilha.
    for linha, p in enumerate(sorted(pedidos, key=lambda x: (x.data_entrega, x.id)), start=2):
        valores = [
            p.id,
            p.status_label,
            p.cliente_nome,
            telefone_formatado(p.cliente_telefone),
            p.data_entrega.strftime("%d/%m/%Y"),
            DIAS[p.data_entrega.weekday()],
            p.modo_label,
            p.endereco or "—",
            p.resumo_itens(),
            round(p.total, 2),
            p.observacoes or "—",
            p.criado_em.strftime("%d/%m/%Y %H:%M"),
        ]
        for col, valor in enumerate(valores, start=1):
            c = ws.cell(row=linha, column=col, value=valor)
            c.border = _BORDA
            c.alignment = Alignment(vertical="top", wrap_text=col in (8, 9, 11))
        ws.cell(row=linha, column=10).number_format = "#,##0.00"
        if p.status == ENTREGUE:
            ws.cell(row=linha, column=2).font = Font(color="7A7A7A", italic=True)

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{ws.max_row}"

    # Linha de total: só o que de fato é dinheiro a entrar ou já entrou.
    ultima = ws.max_row + 1
    ws.cell(row=ultima, column=9, value="TOTAL").font = Font(bold=True)
    soma = sum(p.total for p in pedidos if p.status in (APROVADO, ENTREGUE))
    total_c = ws.cell(row=ultima, column=10, value=round(soma, 2))
    total_c.font = Font(bold=True)
    total_c.number_format = "#,##0.00"


def gerar_planilha(pedidos: list[Pedido], destino=None):
    """Escreve a planilha e devolve o caminho do arquivo gerado."""
    destino = destino or PLANILHA_PADRAO

    aprovados = [p for p in pedidos if p.status == APROVADO]
    sem_recusados = [p for p in pedidos if p.status != RECUSADO]

    wb = Workbook()

    aba_aprovados = wb.active
    aba_aprovados.title = "Aprovados"
    aba_aprovados.sheet_properties.tabColor = "1E7A46"
    _escrever_aba(aba_aprovados, aprovados, "Nenhum pedido aprovado no momento.")

    aba_todos = wb.create_sheet("Todos")
    aba_todos.sheet_properties.tabColor = "17346B"
    _escrever_aba(aba_todos, sem_recusados, "Nenhum pedido registrado ainda.")

    wb.save(destino)
    return destino
